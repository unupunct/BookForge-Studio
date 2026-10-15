"""Manuscript in, finished book out: minimal intervention after import."""
import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.domain.content_blocks import ParagraphBlock, PoetryBlock
from app.importers.base import ImportOptions
from app.layout.model import Glyphs
from app.layout.pagination import LayoutEngine
from app.services.export_service import export_pdf
from app.services.import_service import run_import
from app.services.project_service import NewBookSpec, create_from_manuscript
from app.services.validation_service import validate_book
from app.ui.dialogs.manuscript_dialog import ManuscriptDialog
from tests.conftest import LOREM, RO
from tests.test_ui import qapp, wait_until, window  # noqa: F401  (fixtures)


def novel_docx(path):
    d = Document()
    d.core_properties.title = "Marea Întoarcere"
    d.core_properties.author = "Ana Popescu"
    for n in (1, 2, 3):
        d.add_heading(f"Capitolul {n}", 1)
        for _ in range(7):
            d.add_paragraph(LOREM * 4 + RO)
        q = d.add_paragraph("Un motto centrat")
        q.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.save(path)


def poetry_docx(path):
    d = Document()
    d.core_properties.title = "Versuri"
    d.core_properties.author = "Poet Anonim"
    d.add_heading("Poeme de vară", 1)
    for title in ("Seară", "Dimineață"):
        d.add_paragraph(title).runs[0].bold = True
        d.add_paragraph("")
        for stanza in (["Cade lumina încet", "peste câmpul tăcut,", "ca un gând care trece"],
                       ["Noaptea vine cu pași", "de pisică pe pietre"]):
            for line in stanza:
                p = d.add_paragraph(line)
                p.paragraph_format.space_after = Pt(0)
            d.add_paragraph("")
        d.add_paragraph("")
        d.add_paragraph("")
    d.save(path)


def test_novel_manuscript_becomes_a_finished_book(tmp_path):
    f = tmp_path / "roman.docx"
    novel_docx(f)
    staged = run_import(f, ImportOptions(update_metadata=True))
    book = create_from_manuscript(NewBookSpec(profile="novel", template="literary_fiction"), staged,
                                  ImportOptions())
    assert book.metadata.title == "Marea Întoarcere" and book.metadata.author == "Ana Popescu"
    kinds = [s.kind for s in book.sections]
    assert kinds[:2] == ["front_cover", "title_page"] and kinds.count("chapter") == 3 and "toc" in kinds
    assert book.settings.template == "literary_fiction" and book.settings.chapter_number_style == "roman"
    res = LayoutEngine().build(book)
    rep = validate_book(book, res)
    assert rep.ok, [i.message for i in rep.errors]
    assert not rep.warnings, [i.message for i in rep.warnings]
    first = [s for s in book.sections if s.kind == "chapter"][0]
    pg = res.pages[res.section_pages[first.id]]
    assert pg.label == "1" and pg.side == "recto"  # template: chapters open on right-hand pages
    shown = [g.text for g in pg.items if isinstance(g, Glyphs)]
    assert "I" in shown and "Capitolul 1" in shown  # template's roman chapter number + title
    centred = next(b for s in book.sections for b in s.blocks if isinstance(b, ParagraphBlock) and b.align == "center")
    assert centred.plain_text() == "Un motto centrat"  # centring kept
    pdf = tmp_path / "roman.pdf"
    out = export_pdf(book, pdf, "digital")
    text = "\n".join(p.get_text() for p in pymupdf.open(str(pdf)))
    assert "Marea Întoarcere" in text and "Ana Popescu" in text and "ĂÂÎȘȚ" in text
    assert out.pages == res.page_count


def test_poetry_manuscript_keeps_stanza_spacing_with_template_style(tmp_path):
    f = tmp_path / "versuri.docx"
    poetry_docx(f)
    staged = run_import(f)
    book = create_from_manuscript(NewBookSpec(profile="poetry", template="poetry_free_verse"), staged,
                                  ImportOptions())
    poems = [b for s in book.sections for b in s.blocks if isinstance(b, PoetryBlock)]
    assert [p.title for p in poems] == ["Seară", "Dimineață"]
    assert poems[0].text.split("\n") == ["Cade lumina încet", "peste câmpul tăcut,", "ca un gând care trece", "",
                                         "Noaptea vine cu pași", "de pisică pe pietre"]
    res = LayoutEngine().build(book)
    assert not [i for i in res.issues if i.severity in ("error", "warning")]
    page = res.pages[res.block_pages[poems[0].id][0]]
    ys = sorted({round(g.y, 1) for g in page.items if isinstance(g, Glyphs) and g.block_id == poems[0].id})[1:]
    step = ys[1] - ys[0]
    assert ys[3] - ys[2] == pytest.approx(step * (1 + book.settings.stanza_gap), rel=0.03)  # template's gap
    assert ys[1] - ys[0] == pytest.approx(ys[2] - ys[1])  # even line spacing inside a stanza
    assert page.label and page.matter == "main"  # numbered in the template's folio style
    dest = tmp_path / "versuri.pdf"
    export_pdf(book, dest, "print", allow_errors=True)
    doc = pymupdf.open(str(dest))
    assert doc.page_count == res.page_count and "peste câmpul tăcut," in "\n".join(p.get_text() for p in doc)


def test_manuscript_dialog_creates_spec_and_parses(window, tmp_path):
    f = tmp_path / "roman.docx"
    novel_docx(f)
    dlg = ManuscriptDialog(window, window.state.templates, str(tmp_path), str(f))
    assert wait_until(lambda: dlg.staged is not None)
    assert dlg.title.text() == "Marea Întoarcere" and dlg.author.text() == "Ana Popescu"
    assert dlg.create_btn.isEnabled() and "3 chapter(s)" in dlg.report.toPlainText()
    dlg.category.setCurrentIndex(dlg.category.findData("poetry"))
    assert dlg.template.currentData() == "minimalist_poetry"  # category picks its template
    dlg.template.setCurrentIndex(dlg.template.findData("poetry_haiku_pages"))
    spec = dlg.spec()
    assert spec.profile == "poetry" and spec.template == "poetry_haiku_pages" and spec.title == "Marea Întoarcere"
    book = create_from_manuscript(spec, dlg.staged, dlg.options(), window.state.templates)
    window.state.install_project(book)
    assert window.state.project.settings.template == "poetry_haiku_pages"
    assert window.state.dirty and window.navigator.tree.topLevelItemCount() >= 5
    dlg.close()


def test_manuscript_dialog_reports_bad_file(window, tmp_path):
    f = tmp_path / "x.docx"
    f.write_bytes(b"not a docx")
    dlg = ManuscriptDialog(window, window.state.templates, "", str(f))
    assert wait_until(lambda: "Import failed" in dlg.report.toPlainText())
    assert not dlg.create_btn.isEnabled()
    dlg.close()
