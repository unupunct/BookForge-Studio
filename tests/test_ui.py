import os
import time
from pathlib import Path

import pymupdf
import pytest
from PySide6.QtCore import QCoreApplication, QEventLoop, Qt
from PySide6.QtGui import QKeyEvent, QTextCursor
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox

from app.application_state import AppState
from app.domain.content_blocks import (HeadingBlock, ImageBlock, ParagraphBlock, PoetryBlock,
                                       TableBlock, Run)
from app.domain.sections import Section
from app.main_window import MainWindow
from app.persistence.project_archive import save_project
from app.services.autosave_service import AutosaveService
from app.services.import_service import commit_import, run_import
from app.services.project_service import NewBookSpec, SettingsStore
from app.ui import theme
from app.ui.dialogs.export_dialog import ExportDialog
from app.ui.dialogs.import_dialog import ImportDialog
from app.ui.dialogs.misc_dialogs import SettingsDialog, ValidationDialog
from app.ui.dialogs.search_dialog import SearchDialog
from app.ui.dialogs.wizard import NewBookWizard
from app.ui.editor.block_codec import LINE_SEP
from tests.conftest import LOREM, RO, make_png


@pytest.fixture(scope="session")
def qapp():
    app = QApplication.instance() or QApplication([])
    theme.apply_theme(app)
    return app


def wait_until(cond, timeout=15.0):
    end = time.time() + timeout
    while time.time() < end:
        QCoreApplication.processEvents(QEventLoop.AllEvents, 25)
        if cond():
            return True
        time.sleep(0.01)
    return cond()


@pytest.fixture
def window(qapp, tmp_path):
    state = AppState(autosave=AutosaveService(tmp_path / "rec"), settings=SettingsStore(tmp_path / "s.json"))
    (tmp_path / "rec").mkdir(exist_ok=True)
    state.settings_store.data.autosave_seconds = 0
    state.apply_autosave_interval()
    win = MainWindow(state)
    win.show()
    state.new_project(NewBookSpec(title="Cartea mea", author="Ion", profile="novel"))
    state.mark_dirty(False)
    wait_until(lambda: state.layout is not None)
    yield win
    state.layout_controller.wait_idle()
    win.state._autosave_timer.stop()
    win._discarded = True
    state.mark_dirty(False)
    win.close()
    QCoreApplication.processEvents()


def chapter(win):
    return next(s for s in win.state.project.sections if s.kind == "chapter")


def type_text(win, text):
    win.editor.edit.setFocus()
    cur = win.editor.edit.textCursor()
    cur.movePosition(QTextCursor.End)
    win.editor.edit.setTextCursor(cur)
    QTest.keyClicks(win.editor.edit, text)


def insert_text(win, text):
    """Insert text like an IME/paste would (QTest.keyClicks cannot send non-ASCII keys)."""
    cur = win.editor.edit.textCursor()
    cur.movePosition(QTextCursor.End)
    cur.insertText(text)
    win.editor.edit.setTextCursor(cur)


# ------------------------------------------------------------------ startup
def test_main_window_initialisation(window):
    names = {a.text().replace("&", "").rstrip("…") for a in window.findChildren(type(window.act_new))}
    for needed in ("New book", "Open", "Save", "Save as", "Import manuscript", "Export", "Undo", "Redo",
                   "Preview", "Validate book", "Settings"):
        assert needed in names, needed
    assert window.navigator.tree.topLevelItemCount() >= 5
    assert window.tabs.count() == 2
    assert "Cartea mea" in window.windowTitle()
    assert window.lbl_dirty.text() == "Saved"
    assert window.properties.count() == 4
    assert window.navigator.current_id() == chapter(window).id


def test_wizard_produces_spec(qapp):
    wiz = NewBookWizard()
    wiz.basics.title.setText("Țara lui Papură Vodă")
    wiz.basics.author.setText("Autor")
    wiz.kind.category.setCurrentIndex(wiz.kind.category.findData("cookbook"))
    wiz.layout_page.cover.setChecked(False)
    spec = wiz.spec()
    assert spec.title == "Țara lui Papură Vodă" and spec.profile == "cookbook"
    assert spec.template == "cookbook" and spec.cover is False
    assert wiz.basics.isComplete()
    wiz.basics.title.setText("  ")
    assert not wiz.basics.isComplete()


def test_project_creation_updates_ui(window):
    window.state.new_project(NewBookSpec(title="Poeme", profile="poetry"))
    assert window.state.project.metadata.title == "Poeme"
    assert window.navigator.tree.topLevelItemCount() >= 4
    assert window.state.dirty
    assert window.state.project.settings.template == "minimalist_poetry"


# ------------------------------------------------------------------ editing
def test_typing_reaches_model_and_counts(window):
    type_text(window, "Salut lume ")
    window.state.flush()
    text = " ".join(b.plain_text() for b in chapter(window).blocks)
    assert "Salut lume" in text
    assert window.state.dirty and "Unsaved" in window.lbl_dirty.text()
    assert wait_until(lambda: "2 words here" in window.lbl_words.text() or "words here" in window.lbl_words.text())


def test_romanian_text_typed_in_editor(window):
    # QTest.keyClicks cannot synthesise non-ASCII keys; insert as an IME/paste would.
    window.editor.edit.textCursor().insertText(RO)
    window.state.flush()
    assert RO in chapter(window).blocks[0].plain_text()


def test_block_kind_and_inline_formatting(window):
    ed = window.editor
    type_text(window, "Titlu")
    ed.set_block_kind("heading", 2)
    window.state.flush()
    blk = chapter(window).blocks[0]
    assert isinstance(blk, HeadingBlock) and blk.level == 2 and blk.plain_text() == "Titlu"
    cur = ed.edit.textCursor()
    cur.movePosition(QTextCursor.End)
    ed.edit.setTextCursor(cur)
    QTest.keyClick(ed.edit, Qt.Key_Return)  # Enter after a heading starts a body paragraph
    QTest.keyClicks(ed.edit, "plain bold")
    cur = ed.edit.textCursor()
    cur.movePosition(QTextCursor.EndOfBlock)
    cur.movePosition(QTextCursor.Left, QTextCursor.KeepAnchor, 4)
    ed.edit.setTextCursor(cur)
    ed.toggle_bold()
    window.state.flush()
    blocks = chapter(window).blocks
    assert isinstance(blocks[1], ParagraphBlock)
    assert [(r.text, r.bold) for r in blocks[1].runs] == [("plain ", False), ("bold", True)]


def test_poetry_line_breaks_preserved_in_editor(window):
    ed = window.editor
    ed.set_block_kind("poetry")
    QTest.keyClicks(ed.edit, "Linia unu")
    QTest.keyClick(ed.edit, Qt.Key_Return)
    QTest.keyClicks(ed.edit, "Linia doi")
    QTest.keyClick(ed.edit, Qt.Key_Return)
    QTest.keyClick(ed.edit, Qt.Key_Return)
    QTest.keyClicks(ed.edit, "Strofa doi")
    window.state.flush()
    poem = chapter(window).blocks[0]
    assert isinstance(poem, PoetryBlock)
    assert poem.text == "Linia unu\nLinia doi\n\nStrofa doi"
    assert len(chapter(window).blocks) == 1


def test_insert_table_and_page_break_blocks(window):
    ed = window.editor
    type_text(window, "Before")
    ed.insert_block(TableBlock(rows=[["a", "b"], ["1", "2"]]))
    window.state.flush()
    kinds = [type(b).__name__ for b in chapter(window).blocks]
    assert kinds[:2] == ["ParagraphBlock", "TableBlock"]
    tbl = chapter(window).blocks[1]
    ed.load_section(chapter(window).id)  # reload from the model: placeholder survives
    window.state.flush()
    assert chapter(window).blocks[1].id == tbl.id and chapter(window).blocks[1].rows == [["a", "b"], ["1", "2"]]


def test_placeholder_is_protected_from_typing(window):
    ed = window.editor
    ed.insert_block(TableBlock(rows=[["x"]]))
    ed.edit.moveCursor(QTextCursor.Start)
    ed.edit.moveCursor(QTextCursor.NextBlock)
    assert ed.current_kind() == "placeholder"
    QTest.keyClicks(ed.edit, "oops")
    window.state.flush()
    assert any(isinstance(b, TableBlock) for b in chapter(window).blocks)
    assert "oops" not in " ".join(b.plain_text() for b in chapter(window).blocks)


def test_locked_block_cannot_be_edited(window):
    ed = window.editor
    type_text(window, "keep me")
    window.state.flush()
    ed.set_locked(True)
    type_text(window, "XXX")
    window.state.flush()
    assert "XXX" not in chapter(window).blocks[0].plain_text()
    ed.set_locked(False)
    type_text(window, "YYY")
    window.state.flush()
    assert "YYY" in chapter(window).blocks[0].plain_text()


def test_paste_inserts_plain_text_only(window):
    from PySide6.QtCore import QMimeData
    md = QMimeData()
    md.setHtml("<h1 style='color:red'>Pasted <b>rich</b></h1><script>x</script>")
    md.setText("Pasted rich")
    window.editor.edit.insertFromMimeData(md)
    window.state.flush()
    blk = chapter(window).blocks[0]
    assert isinstance(blk, ParagraphBlock) and blk.plain_text() == "Pasted rich"
    assert all(not r.bold for r in blk.runs)


def test_editor_find_and_replace(window):
    type_text(window, "aaa bbb aaa")
    ed = window.editor
    ed.show_find()
    ed.find_bar.find.setText("aaa")
    ed.find_bar.repl.setText("zzz")
    ed.find_bar.replace_all()
    window.state.flush()
    assert chapter(window).blocks[0].plain_text() == "zzz bbb zzz"
    ed.edit.undo()
    window.state.flush()
    assert chapter(window).blocks[0].plain_text() == "aaa bbb aaa"


# -------------------------------------------------------- navigator/structure
def test_structure_add_rename_reorder_undo(window):
    s = window.state
    n = len(list(s.project.walk()))
    new = Section("Capitolul nou", "chapter")
    s.add_section(new)
    assert len(list(s.project.walk())) == n + 1
    window.navigator.rebuild()
    window.navigator.select(new.id)
    s.rename_section(new.id, "Redenumit")
    assert s.project.find_section(new.id).title == "Redenumit"
    assert s.move_section(new.id, -1)
    s.undo_stack.undo()  # move
    s.undo_stack.undo()  # rename
    assert s.project.find_section(new.id).title == "Capitolul nou"
    s.undo_stack.undo()  # add
    assert len(list(s.project.walk())) == n
    s.undo_stack.redo()
    assert s.project.find_section(new.id) is not None


def test_drag_and_drop_reorder_through_state(window):
    s = window.state
    p = s.project
    ch = chapter(window)
    extra = Section("Second", "chapter")
    s.add_section(extra)
    top = [(x.id, []) for x in p.sections]
    ids = [i for i, _ in top]
    a, b = ids.index(ch.id), ids.index(extra.id)
    top[a], top[b] = top[b], top[a]
    s.reorder_from_tree(top)
    assert [x.id for x in p.sections].index(extra.id) < [x.id for x in p.sections].index(ch.id)
    top2 = [(x.id, [(extra.id, [])] if x.id == ch.id else []) for x in p.sections if x.id != extra.id]
    s.reorder_from_tree(top2)
    assert [c.id for c in ch.children] == [extra.id]


def test_navigator_hide_and_remove_with_confirmation(window, monkeypatch):
    nav, s = window.navigator, window.state
    sec = Section("Temp", "chapter", blocks=[ParagraphBlock(runs=[Run("data")])])
    s.add_section(sec)
    nav.rebuild()
    nav.select(sec.id)
    monkeypatch.setattr(QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.No))
    nav.delete_current()
    assert s.project.find_section(sec.id) is not None
    monkeypatch.setattr(QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.Yes))
    nav.delete_current()
    assert s.project.find_section(sec.id) is None
    s.undo_stack.undo()
    assert s.project.find_section(sec.id).blocks[0].plain_text() == "data"
    nav.rebuild()
    nav.select(sec.id)
    nav.toggle_enabled()
    assert s.project.find_section(sec.id).enabled is False


# ----------------------------------------------------------------- import
def test_import_dialog_workflow(window, tmp_path):
    f = tmp_path / "manuscris.txt"
    f.write_text("Povestea mea\n\nCapitolul 1\n\nProlog cu diacritice: ăâîșț.\n\nCapitolul 2\n\nFinal.", encoding="utf-8")
    dlg = ImportDialog(window, str(tmp_path), str(f))
    assert wait_until(lambda: dlg.staged is not None)
    assert "2 chapter(s)" in dlg.report.toPlainText() and dlg.add_btn.isEnabled()
    n = len(window.state.project.sections)
    window.state.edit("Import", lambda p: commit_import(p, dlg.staged, dlg.options))
    window.navigator.rebuild()
    assert len(window.state.project.sections) == n + 2
    assert window.state.project.sources["manuscris.txt"]
    dlg.close()


def test_import_dialog_failure_is_reported(window, tmp_path):
    f = tmp_path / "x.xyz"
    f.write_text("nope")
    dlg = ImportDialog(window, "", str(f))
    assert wait_until(lambda: "Import failed" in dlg.report.toPlainText())
    assert not dlg.add_btn.isEnabled() and dlg.staged is None
    dlg.close()


def test_import_docx_into_window_and_preview(window, tmp_path):
    from docx import Document
    d = Document()
    d.add_heading("Capitolul Unu", 1)
    d.add_paragraph("Un paragraf cu diacritice: " + RO)
    f = tmp_path / "m.docx"
    d.save(f)
    staged = run_import(f)
    window.state.edit("Import", lambda p: commit_import(p, staged))
    window.state.request_layout(immediate=True)
    window.state.layout_controller.wait_idle()
    assert wait_until(lambda: window.state.layout is not None and any(
        "Capitolul Unu" == s.title for s in window.state.project.sections))
    assert any(r.title == "Capitolul Unu" for r in [s for s in window.state.project.sections])


# ------------------------------------------------------------ save and open
def test_save_and_open_workflow(window, tmp_path):
    type_text(window, "Text salvat")
    p = tmp_path / "carte.bookforge"
    window.state.save(p)
    assert p.exists() and not window.state.dirty and not (tmp_path / "carte.bookforge.tmp").exists()
    other = AppState(autosave=AutosaveService(tmp_path / "rec2"), settings=SettingsStore(tmp_path / "s2.json"))
    other.open_project(p)
    assert other.project.metadata.title == "Cartea mea"
    assert "Text salvat" in " ".join(b.plain_text() for s in other.project.sections for b in s.blocks)
    assert str(p.resolve()) in other.settings_store.existing_recent()


def test_open_corrupt_file_reports_error_and_keeps_project(window, tmp_path, monkeypatch):
    bad = tmp_path / "bad.bookforge"
    bad.write_bytes(b"garbage")
    errors = []
    monkeypatch.setattr(window, "show_error", lambda title, msg, details="": errors.append((title, msg)))
    before = window.state.project.metadata.title
    assert window._open_path(bad) is False
    assert errors and "could not be opened" in errors[0][1]
    assert window.state.project.metadata.title == before


def test_open_corrupt_file_offers_backup(window, tmp_path, monkeypatch):
    good = tmp_path / "g.bookforge"
    window.state.project.metadata.title = "Versiunea veche"
    window.state.save(good)
    window.state.project.metadata.title = "Versiunea noua"
    window.state.save(good)
    good.write_bytes(b"corrupted after the fact")
    monkeypatch.setattr(QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.Yes))
    assert window._open_path(good) is True
    assert window.state.project.metadata.title == "Versiunea veche"
    assert window.state.dirty  # recovered content must be saved explicitly


def test_save_failure_is_reported(window, tmp_path, monkeypatch):
    errors = []
    monkeypatch.setattr(window, "show_error", lambda t, m, d="": errors.append(m))
    window.state.path = tmp_path / "no_such_dir" / "x.bookforge"
    (tmp_path / "blocker").write_text("file, not a directory")
    window.state.path = tmp_path / "blocker" / "x.bookforge"
    assert window.save() is False
    assert errors and "could not be saved" in errors[0]


def test_close_with_unsaved_changes_prompts(window, monkeypatch):
    type_text(window, "x")
    window.state.flush()
    monkeypatch.setattr(QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.Cancel))
    window._discarded = False
    assert window.confirm_discard() is False


# ------------------------------------------------------------------ preview
def test_preview_matches_layout_and_navigates(window):
    s = window.state
    insert_text(window, LOREM * 40)
    s.flush()
    s.request_layout(immediate=True)
    s.layout_controller.wait_idle()
    assert wait_until(lambda: window.preview.layout_result is not None and window.preview.layout_result.page_count >= 6)
    pv = window.preview
    assert pv.layout_result is s.layout
    assert pv.thumbs.count() == s.layout.page_count
    ch = chapter(window)
    pv.show_section(ch.id)
    assert pv.current == s.layout.section_pages[ch.id]
    img = pv.page_image(s.layout.pages[pv.current], 1.0)
    assert not img.isNull() and img.width() == round(s.layout.page_width)
    # the page really contains painted content (not blank paper)
    colors = {img.pixel(x, y) for x in range(0, img.width(), 7) for y in range(0, img.height(), 7)}
    assert len(colors) > 3
    pv._mode_changed(1)
    assert len(pv.visible_pages()) in (1, 2)
    pv.go_relative(1)
    pv.show_block(ch.blocks[0].id)
    assert pv.highlight_block == ch.blocks[0].id


def test_preview_flags_pages_with_issues(window):
    s = window.state
    ch = chapter(window)
    ch.blocks.append(TableBlock(rows=[["h"], [" ".join(["word"] * 4000)]]))
    s.request_layout(immediate=True)
    s.layout_controller.wait_idle()
    assert wait_until(lambda: window.preview.layout_result is not None and window.preview._issues_by_page)
    assert window.preview.issue_btn.isEnabled()
    page_idx = next(iter(window.preview._issues_by_page))
    window.preview.next_issue()
    assert window.preview.current == page_idx


# ------------------------------------------------------------------ export
def run_export(dlg, timeout=30):
    dlg.go.click()
    assert wait_until(lambda: dlg._stage in ("review", "idle") and dlg.go.isEnabled(), timeout)
    if dlg._stage == "review":
        dlg.go.click()
        assert wait_until(lambda: dlg._stage == "idle" and dlg.go.isEnabled() and
                          (dlg.result is not None or "✖" in dlg.status.text()), timeout)


def test_export_dialog_end_to_end_pdf(window, tmp_path):
    insert_text(window, "Un paragraf pentru export. " + RO)
    dest = tmp_path / "carte.pdf"
    dlg = ExportDialog(window.state, window)
    dlg.path.setText(str(dest))
    run_export(dlg)
    assert dlg.result is not None, dlg.status.text()
    doc = pymupdf.open(str(dest))
    assert doc.page_count == dlg.result.pages
    assert "Un paragraf pentru export" in "\n".join(p.get_text() for p in doc)
    assert doc.metadata["title"] == "Cartea mea"
    dlg.close()


def test_export_dialog_shows_errors_and_requires_confirmation(window, tmp_path):
    ch = chapter(window)
    ch.blocks.append(ImageBlock(asset="missing-asset"))
    dest = tmp_path / "bad.pdf"
    dlg = ExportDialog(window.state, window)
    dlg.path.setText(str(dest))
    dlg.go.click()
    assert wait_until(lambda: dlg._stage == "review", 30)
    assert dlg.issues.topLevelItemCount() >= 1
    assert dlg.issues.topLevelItem(0).text(0) == "ERROR" and not dlg.force.isHidden()
    dlg.go.click()  # without confirming the errors nothing is exported
    QCoreApplication.processEvents()
    assert not dest.exists() and "despite" in dlg.status.text()
    dlg.force.setChecked(True)
    dlg.go.click()
    assert wait_until(lambda: dest.exists() and dlg.result is not None, 30)
    dlg.close()


def test_export_never_overwrites_without_consent(window, tmp_path, monkeypatch):
    dest = tmp_path / "exists.pdf"
    dest.write_bytes(b"old")
    monkeypatch.setattr(QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.No))
    dlg = ExportDialog(window.state, window)
    dlg.path.setText(str(dest))
    dlg.go.click()
    QCoreApplication.processEvents()
    assert dest.read_bytes() == b"old" and dlg._stage == "idle"
    dlg.close()


def test_export_other_formats_via_dialog(window, tmp_path):
    for key, name in (("html", "b.html"), ("markdown", "b.md"), ("docx", "b.docx")):
        dlg = ExportDialog(window.state, window, initial_format=key)
        dlg.path.setText(str(tmp_path / name))
        run_export(dlg)
        assert (tmp_path / name).exists(), (key, dlg.status.text())
        dlg.close()


# ------------------------------------------------------------- misc dialogs
def test_validation_dialog_navigation(window):
    ch = chapter(window)
    ch.blocks.append(ImageBlock(asset="ghost"))
    window.validate()
    assert wait_until(lambda: window.last_report is not None, 30)
    assert not window.last_report.ok
    assert "error" in window.btn_valid.text()
    dlg = window.validation_dialog
    issue = next(i for i in window.last_report.issues if i.code == "missing_asset")
    window._navigate_issue(issue)
    assert window.navigator.current_id() == ch.id
    dlg.close()


def test_search_dialog_book_wide(window, monkeypatch):
    s = window.state
    ch = chapter(window)
    ch.blocks = [ParagraphBlock(runs=[Run("pisica si pisica")]),
                 HeadingBlock(level=1, runs=[Run("Pisica")])]
    window.editor.load_section(ch.id)
    s.add_section(Section("Doi", "chapter", blocks=[ParagraphBlock(runs=[Run("o pisica mare")])]))
    dlg = SearchDialog(s, window)
    dlg.find.setText("pisica")
    dlg.run_search()
    assert len(dlg.hits) == 4
    monkeypatch.setattr(QMessageBox, "question", staticmethod(lambda *a, **k: QMessageBox.Yes))
    dlg.repl.setText("catel")
    replaced = []
    dlg.replaced.connect(lambda: replaced.append(1))
    dlg.replace_everything()
    assert replaced and not search_hits(s, "pisica")
    window._after_replace()
    assert "catel" in window.editor.edit.toPlainText()
    s.undo_stack.undo()
    assert search_hits(s, "pisica")


def search_hits(state, q):
    from app.services.search_service import search_book
    return search_book(state.project, q)


def test_settings_dialog_applies_autosave_interval(window):
    dlg = SettingsDialog(window.state, window)
    dlg.autosave.setValue(5)
    dlg._ok()
    assert window.state.settings_store.data.autosave_seconds == 5
    assert window.state._autosave_timer.isActive()
    window.state.settings_store.data.autosave_seconds = 0
    window.state.apply_autosave_interval()


def test_autosave_writes_recovery_file_without_blocking(window, tmp_path):
    s = window.state
    type_text(window, "nesalvat")
    s.flush()
    t = time.perf_counter()
    s._autosave_tick()
    assert time.perf_counter() - t < 1.0  # snapshot only; the write happens in the background
    assert s.autosave.wait_idle()
    cands = s.autosave.candidates()
    assert cands and cands[0].project_id == s.project.id
    assert "nesalvat" not in cands[0].title
    recovered = s.autosave.load(cands[0])
    assert "nesalvat" in " ".join(b.plain_text() for sec in recovered.sections for b in sec.blocks)
    s.open_recovered(cands[0])
    assert s.dirty


# --------------------------------------------------------- properties panel
def test_book_properties_change_layout_and_undo(window):
    bt = window.properties.book_tab
    s = window.state
    old = s.project.settings.margins.inner
    bt.m_in.setValue(30.0)
    assert s.project.settings.margins.inner == 30.0
    s.undo_stack.undo()
    assert s.project.settings.margins.inner == old
    bt.size.setCurrentText("A4")
    bt._page_size()
    assert s.project.settings.page_size == "A4" and abs(s.project.settings.width_mm - 210) < 0.01
    s.layout_controller.wait_idle()
    wait_until(lambda: s.layout and abs(s.layout.page_width - 595.28) < 1)
    assert abs(s.layout.page_width - 595.28) < 1


def test_style_edit_marks_customised_and_survives_template(window):
    st = window.properties.styles_tab
    s = window.state
    st.name.setCurrentIndex(st.name.findData("heading1"))
    st.load_style()
    st.size.setValue(26)
    assert s.project.styles["heading1"].font_size == 26 and s.project.styles["heading1"].user_modified
    bt = window.properties.book_tab
    from app.services.template_service import apply_template
    apply_template(s.project, s.templates.get("technical_handbook"))
    assert s.project.styles["heading1"].font_size == 26


def test_typography_setting_changes_page_count(window):
    s = window.state
    ch = chapter(window)
    ch.blocks = [ParagraphBlock(runs=[Run(LOREM * 40)]) for _ in range(6)]
    s.request_layout(immediate=True)
    s.layout_controller.wait_idle()
    wait_until(lambda: s.layout is not None)
    before = s.layout.page_count
    window.properties.book_tab._set("body_size", 16.0)
    s.layout_controller.wait_idle()
    assert wait_until(lambda: s.layout.page_count > before)


def test_section_tab_edits(window):
    tab = window.properties.section_tab
    ch = chapter(window)
    window.navigator.select(ch.id)
    tab.title.setText("Titlu nou")
    tab._title_done()
    assert ch.title == "Titlu nou" or window.state.project.find_section(ch.id).title == "Titlu nou"
    tab.in_toc.setChecked(False)
    assert window.state.project.find_section(ch.id).in_toc is False
    tab.new_page.setCurrentIndex(2)
    tab._apply(new_page=tab.new_page.currentData())
    assert window.state.project.find_section(ch.id).new_page is False
