import io

import pytest
from docx import Document
from docx.shared import Inches

from app.domain.book import BookProject
from app.domain.content_blocks import (CodeBlock, HeadingBlock, ImageBlock, ListBlock,
                                       PageBreakBlock, ParagraphBlock, PoetryBlock, QuoteBlock,
                                       SceneBreakBlock, TableBlock, FootnoteBlock, runs_text)
from app.importers.base import ImportFailure, ImportOptions
from app.services.import_service import commit_import, import_file, run_import
from tests.conftest import RO, make_png


# ------------------------------------------------------------------- TXT
def test_txt_utf8_romanian_and_chapters(tmp_path):
    f = tmp_path / "carte.txt"
    f.write_text(f"Titlul cărții\n\nCapitolul 1\n\nPrima linie {RO}\ncontinuă aici.\n\n"
                 f"Al doilea paragraf.\n\nCapitolul 2\n\nAltceva.\n", encoding="utf-8")
    res = run_import(f).result
    assert res.title == "Titlul cărții"
    assert [s.title for s in res.sections] == ["Capitolul 1", "Capitolul 2"]
    first = res.sections[0].blocks[0]
    assert "continuă aici." in first.plain_text() and RO in first.plain_text()
    assert len(res.sections[0].blocks) == 2


def test_txt_bom_and_fallback_encodings(tmp_path):
    f = tmp_path / "bom.txt"
    f.write_bytes("﻿Școală\n\nText".encode("utf-8"))
    res = run_import(f, ImportOptions(detect_title=False)).result
    assert res.sections[0].blocks[0].plain_text() == "Școală"
    assert any("byte-order mark" in i.message for i in res.report.issues)
    g = tmp_path / "legacy.txt"
    g.write_bytes("café\n\nfoo".encode("cp1252"))
    res2 = run_import(g, ImportOptions(detect_title=False)).result
    assert res2.sections[0].blocks[0].plain_text() == "café"
    assert res2.report.has_warnings
    h = tmp_path / "latin2.txt"
    h.write_bytes("Ștefan ț".encode("cp1250", errors="replace"))
    res3 = run_import(h, ImportOptions(encoding="cp1250", detect_title=False)).result
    assert "tefan" in res3.sections[0].blocks[0].plain_text()


def test_txt_unknown_encoding_and_custom_pattern(tmp_path):
    f = tmp_path / "x.txt"
    f.write_text("hello", encoding="utf-8")
    with pytest.raises(ImportFailure):
        run_import(f, ImportOptions(encoding="nope-1"))
    g = tmp_path / "y.txt"
    g.write_text("=== ONE ===\n\ntext\n\n=== TWO ===\n\nmore\n", encoding="utf-8")
    res = run_import(g, ImportOptions(chapter_pattern=r"^===.*===$", detect_title=False)).result
    assert [s.title for s in res.sections] == ["=== ONE ===", "=== TWO ==="]
    with pytest.raises(ImportFailure):
        run_import(g, ImportOptions(chapter_pattern="("))


def test_txt_poetry_mode_keeps_lines_and_scene_breaks(tmp_path):
    f = tmp_path / "p.txt"
    f.write_text("Line one\nLine two\n\nStanza two\n\n\nProse here\n\n* * *\n\nAfter", encoding="utf-8")
    res = run_import(f, ImportOptions(preserve_line_breaks=True, detect_title=False)).result
    blocks = res.sections[0].blocks
    assert isinstance(blocks[0], PoetryBlock) and blocks[0].text == "Line one\nLine two\n\nStanza two"
    assert any(isinstance(b, SceneBreakBlock) for b in blocks)


def test_import_never_modifies_source_and_keeps_snapshot(tmp_path):
    f = tmp_path / "orig.txt"
    f.write_text("Alpha\n\nBeta", encoding="utf-8")
    before = f.read_bytes()
    mtime = f.stat().st_mtime_ns
    book = BookProject()
    import_file(book, f)
    assert f.read_bytes() == before and f.stat().st_mtime_ns == mtime
    assert book.sources["orig.txt"] == before
    assert book.import_reports and book.import_reports[0]["format"] == "TXT"


def test_unsupported_and_rtf(tmp_path):
    f = tmp_path / "a.xyz"
    f.write_text("x")
    with pytest.raises(ImportFailure, match="Unsupported"):
        run_import(f)
    r = tmp_path / "a.rtf"
    r.write_text("{\\rtf1 hi}")
    with pytest.raises(ImportFailure, match="RTF"):
        run_import(r)
    with pytest.raises(ImportFailure):
        run_import(tmp_path / "missing.txt")


# -------------------------------------------------------------- Markdown
MD = """# Cartea mea

Intro with **bold**, *italic*, `code` and a [link](https://example.com).

## Capitolul unu

Primul paragraf cu diacritice: Ștefan și țăranii.

- item unu
- item doi
  - nested

1. first
2. second

> A quote

```python
def f():
    return "ș"
```

| A | B |
|---|---|
| 1 | 2 |

---

![alt text](pic.png "A caption")
![remote](https://example.com/x.png)

<div>raw html</div>

## Capitolul doi

Text doi.
"""


def test_markdown_structures(tmp_path):
    (tmp_path / "pic.png").write_bytes(make_png())
    f = tmp_path / "doc.md"
    f.write_text(MD, encoding="utf-8")
    book = BookProject()
    res = import_file(book, f, ImportOptions(split_level=2))
    titles = [s.title for s in res.sections]
    assert titles == ["Doc", "Capitolul unu", "Capitolul doi"] or titles[1:] == ["Capitolul unu", "Capitolul doi"]
    chap = next(s for s in res.sections if s.title == "Capitolul unu")
    kinds = [type(b).__name__ for b in chap.blocks]
    for k in ("ParagraphBlock", "ListBlock", "QuoteBlock", "CodeBlock", "TableBlock",
              "SceneBreakBlock", "ImageBlock"):
        assert k in kinds, kinds
    lists = [b for b in chap.blocks if isinstance(b, ListBlock)]
    assert lists[0].ordered is False and len(lists[0].items) == 3 and lists[1].ordered
    code = next(b for b in chap.blocks if isinstance(b, CodeBlock))
    assert 'return "ș"' in code.text and code.language == "python"
    table = next(b for b in chap.blocks if isinstance(b, TableBlock))
    assert table.rows == [["A", "B"], ["1", "2"]] and table.header_rows == 1
    img = next(b for b in chap.blocks if isinstance(b, ImageBlock))
    assert runs_text(img.caption) == "A caption" and img.asset in book.assets
    msgs = " ".join(i.message for i in res.report.issues)
    assert "Remote image" in msgs and "HTML" in msgs and "flattened" in msgs
    intro = next(s for s in res.sections if s.title != "Capitolul unu")
    para = next(b for b in intro.blocks if isinstance(b, ParagraphBlock))
    assert any(r.bold and r.text == "bold" for r in para.runs)
    assert any(r.italic for r in para.runs) and any(r.code for r in para.runs)
    assert any(r.link == "https://example.com" for r in para.runs)


def test_markdown_image_cannot_escape_folder(tmp_path):
    secret = tmp_path / "secret.png"
    secret.write_bytes(make_png())
    sub = tmp_path / "sub"
    sub.mkdir()
    f = sub / "d.md"
    f.write_text("![x](../secret.png)\n\ntext", encoding="utf-8")
    book = BookProject()
    res = import_file(book, f)
    assert not book.assets
    assert any("outside" in i.message for i in res.report.issues)


# ------------------------------------------------------------------ HTML
def test_html_import_ignores_scripts(tmp_path):
    f = tmp_path / "p.html"
    f.write_text("""<html><head><title>Titlu</title><script>alert(1)</script></head><body>
    <h1>Cap 1</h1><p>Hello <b>bold</b> și <a href="http://a.b">link</a>.</p>
    <ul><li>one</li><li>two</li></ul><pre>a  b
 c</pre><table><tr><th>x</th></tr><tr><td>1</td></tr></table>
    <script>evil()</script></body></html>""", encoding="utf-8")
    res = run_import(f).result
    chap = res.sections[0]
    assert chap.title == "Cap 1"
    assert any(isinstance(b, ListBlock) for b in chap.blocks)
    assert any(isinstance(b, CodeBlock) and "a  b" in b.text for b in chap.blocks)
    assert any(isinstance(b, TableBlock) for b in chap.blocks)
    assert "evil" not in " ".join(b.plain_text() for b in chap.blocks)
    assert any("script" in i.message for i in res.report.issues)


# ------------------------------------------------------------------ DOCX
def _docx(path, png):
    d = Document()
    d.core_properties.title = "Docx Title"
    d.add_heading("Capitolul Întâi", 1)
    p = d.add_paragraph("Text ")
    p.add_run("gras").bold = True
    p.add_run(" și ")
    p.add_run("cursiv").italic = True
    p.add_run(f" {RO}")
    d.add_heading("Subsecțiune", 2)
    d.add_paragraph("item", style="List Bullet")
    d.add_paragraph("item2", style="List Bullet")
    d.add_paragraph("pas", style="List Number")
    d.add_picture(io.BytesIO(png), width=Inches(2))
    d.add_paragraph("Figura 1: test", style="Caption")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "H1", "H2", "a", "b"
    d.add_page_break()
    d.add_heading("Capitolul Doi", 1)
    d.add_paragraph("Al doilea.")
    d.save(path)


def test_docx_headings_formatting_lists_tables_images(tmp_path, png_bytes):
    f = tmp_path / "m.docx"
    _docx(f, png_bytes)
    book = BookProject()
    res = import_file(book, f)
    assert [s.title for s in res.sections] == ["Capitolul Întâi", "Capitolul Doi"]
    ch = res.sections[0]
    para = next(b for b in ch.blocks if isinstance(b, ParagraphBlock))
    assert any(r.bold and r.text == "gras" for r in para.runs)
    assert any(r.italic and r.text == "cursiv" for r in para.runs)
    assert RO in para.plain_text()
    heads = [b for b in ch.blocks if isinstance(b, HeadingBlock)]
    assert heads and heads[0].level == 1 and runs_text(heads[0].runs) == "Subsecțiune"
    lists = [b for b in ch.blocks if isinstance(b, ListBlock)]
    assert lists[0].ordered is False and len(lists[0].items) == 2 and lists[1].ordered is True
    img = next(b for b in ch.blocks if isinstance(b, ImageBlock))
    assert book.assets[img.asset].width == 200 and runs_text(img.caption) == "Figura 1: test"
    tbl = next(b for b in ch.blocks if isinstance(b, TableBlock))
    assert tbl.rows == [["H1", "H2"], ["a", "b"]]
    assert any(isinstance(b, PageBreakBlock) for b in ch.blocks)
    assert book.sources["m.docx"] == f.read_bytes()


def test_docx_hyperlink_footnotes_and_report(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    d = Document()
    d.add_heading("Chap", 1)
    p = d.add_paragraph("See ")
    r_id = d.part.relate_to("https://example.org", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "the site"
    run.append(t)
    h.append(run)
    p._p.append(h)
    f = tmp_path / "h.docx"
    d.save(f)
    res = run_import(f).result
    para = next(b for b in res.sections[0].blocks if isinstance(b, ParagraphBlock))
    assert any(r.link == "https://example.org" and r.text == "the site" for r in para.runs)


def test_docx_corrupt_file_reports_failure(tmp_path):
    f = tmp_path / "bad.docx"
    f.write_bytes(b"PK\x03\x04 not really a docx")
    with pytest.raises(ImportFailure):
        run_import(f)
    g = tmp_path / "empty.docx"
    g.write_bytes(b"")
    with pytest.raises(ImportFailure):
        run_import(g)


def test_docx_merged_cells_reported(tmp_path):
    d = Document()
    d.add_heading("C", 1)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).merge(t.cell(0, 1))
    t.cell(0, 0).text = "merged"
    f = tmp_path / "mc.docx"
    d.save(f)
    res = run_import(f).result
    assert any("Merged" in i.message for i in res.report.issues)


# ------------------------------------------------------------------- PDF
def test_pdf_text_and_scanned_pages(tmp_path, png_bytes):
    import pymupdf
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "A page with real extractable text for the importer to read.")
    scanned = doc.new_page()
    scanned.insert_image(pymupdf.Rect(50, 50, 300, 300), stream=png_bytes)
    f = tmp_path / "ref.pdf"
    doc.save(str(f))
    res = run_import(f).result
    text = " ".join(b.plain_text() for b in res.sections[0].blocks)
    assert "extractable text" in text
    assert any("scanned" in i.message for i in res.report.issues)
    assert "no extractable text" in text  # placeholder: nothing silently dropped


def test_commit_dedupes_assets(tmp_path, png_bytes):
    (tmp_path / "pic.png").write_bytes(png_bytes)
    f = tmp_path / "a.md"
    f.write_text("![a](pic.png)\n\n![b](pic.png)", encoding="utf-8")
    book = BookProject()
    import_file(book, f)
    assert len(book.assets) == 1
