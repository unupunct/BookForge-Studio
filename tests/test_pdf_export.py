import os
import stat
from pathlib import Path

import pymupdf
import pytest

from app.domain.content_blocks import (CodeBlock, ImageBlock, ParagraphBlock, PoetryBlock, Run,
                                       TableBlock)
from app.domain.sections import Section
from app.domain.layouts import MM_TO_PT
from app.layout.model import Glyphs, ImageItem
from app.layout.pagination import LayoutEngine
from app.services.export_service import (ExportError, ValidationBlocked, export_other,
                                         export_pdf)
from tests.conftest import LOREM, RO, make_png


@pytest.fixture
def exported(tmp_path, sample_book):
    dest = tmp_path / "out.pdf"
    res = export_pdf(sample_book, dest, "digital")
    return sample_book, dest, res


def test_valid_pdf_with_metadata_and_page_count(exported):
    book, dest, res = exported
    assert dest.read_bytes()[:5] == b"%PDF-"
    doc = pymupdf.open(str(dest))
    assert doc.page_count == res.pages == LayoutEngine().build(book).page_count
    assert doc.metadata["title"] == "Cartea de test"
    assert doc.metadata["author"] == "Autor Țărănescu"
    assert not list(dest.parent.glob("*.tmp"))


def test_page_dimensions_match_settings(exported):
    book, dest, _ = exported
    w, h = book.settings.page_size_pt()
    doc = pymupdf.open(str(dest))
    for page in doc:
        assert page.rect.width == pytest.approx(w, abs=0.5)
        assert page.rect.height == pytest.approx(h, abs=0.5)


def test_text_extraction_and_romanian_diacritics(exported):
    _, dest, _ = exported
    doc = pymupdf.open(str(dest))
    text = "\n".join(p.get_text() for p in doc)
    for frag in ("Ștefan își ține ceaiul în căni țesute", "ĂÂÎȘȚ ăâîșț", "„ghilimele” — …",
                 "Capitolul 1", "Roses are red", "Autor Țărănescu"):
        assert frag in text, frag


def test_fonts_embedded(exported):
    _, dest, res = exported
    assert res.preflight["all_fonts_embedded"]
    doc = pymupdf.open(str(dest))
    fonts = {f[3]: f for p in doc for f in p.get_fonts()}
    assert fonts and all(f[1] == "ttf" for f in fonts.values()), fonts


def test_bookmarks_and_internal_links(exported):
    book, dest, _ = exported
    doc = pymupdf.open(str(dest))
    toc = doc.get_toc()
    titles = [t[1] for t in toc]
    assert "Capitolul 1" in titles and "Secțiunea 1" in titles
    levels = [t[0] for t in toc]
    assert levels[0] == 1 and all(b - a <= 1 for a, b in zip(levels, levels[1:]))
    layout = LayoutEngine().build(book)
    ch1 = book.sections[2]
    target = layout.section_pages[ch1.id] + 1
    assert [t for t in toc if t[1] == "Capitolul 1"][0][2] == target
    toc_page = layout.section_pages[book.sections[1].id]
    links = doc[toc_page].get_links()
    assert links and all(l["kind"] == pymupdf.LINK_GOTO for l in links)
    assert any(l["page"] == target - 1 for l in links)


def test_page_numbers_printed_and_match_layout(exported):
    book, dest, _ = exported
    layout = LayoutEngine().build(book)
    doc = pymupdf.open(str(dest))
    for p in layout.pages:
        if not p.label:
            continue
        words = [w[4] for w in doc[p.index].get_text("words")]
        assert p.label in words, (p.index, p.label)


def test_pdf_text_positions_match_layout_exactly(exported):
    """Preview and PDF draw the same layout model: positions must agree."""
    book, dest, _ = exported
    layout = LayoutEngine().build(book)
    doc = pymupdf.open(str(dest))
    checked = 0
    for p in layout.pages[:12]:
        pdf_words = doc[p.index].get_text("words")
        for it in p.items:
            if isinstance(it, Glyphs) and it.text.strip() and not it.word_space:
                first = it.text.split()[0]
                cands = [w for w in pdf_words if w[4] == first and abs(w[0] - it.x) < 1.5
                         and abs(w[3] - it.y) < it.size]
                assert cands, (p.index, first, it.x, it.y)
                checked += 1
    assert checked > 30


def test_image_placement_in_pdf(tmp_path, png_bytes):
    from app.domain.book import BookProject
    b = BookProject()
    b.metadata.title = "Img"
    a = b.add_asset(make_png(600, 300, (10, 200, 30)), "i.png")
    b.sections = [Section("C", "chapter", blocks=[
        ParagraphBlock(runs=[Run("Intro")]), ImageBlock(asset=a.id, width_pct=50, alt="x",
                                                        caption=[Run("Cap")])])]
    layout = LayoutEngine().build(b)
    dest = tmp_path / "i.pdf"
    export_pdf(b, dest, "digital")
    it = next(i for p in layout.pages for i in p.items if isinstance(i, ImageItem))
    page_idx = layout.block_pages[it.block_id][0]
    info = pymupdf.open(str(dest))[page_idx].get_image_info()
    assert len(info) == 1
    x0, y0, x1, y1 = info[0]["bbox"]
    assert (x0, y0, x1 - x0, y1 - y0) == pytest.approx((it.x, it.y, it.w, it.h), abs=0.6)


def test_poetry_and_code_in_pdf_text(tmp_path):
    from app.domain.book import BookProject
    b = BookProject()
    b.metadata.title = "P"
    b.sections = [Section("C", "chapter", blocks=[
        PoetryBlock(text="First line\nSecond line\n\nThird line"),
        CodeBlock(text="print('ș')\n    indented()")])]
    dest = tmp_path / "p.pdf"
    export_pdf(b, dest, "digital")
    text = pymupdf.open(str(dest))[0].get_text()
    lines = [l.strip() for l in text.splitlines()]
    assert lines.index("First line") + 1 == lines.index("Second line")
    assert "Third line" in lines and "print('ș')" in lines


def test_print_profile_bleed_and_crop_marks(tmp_path, sample_book):
    dest = tmp_path / "print.pdf"
    res = export_pdf(sample_book, dest, "print", crop_marks=True)
    w, h = sample_book.settings.page_size_pt()
    ofs = 3 * MM_TO_PT + 9 * MM_TO_PT
    doc = pymupdf.open(str(dest))
    assert doc[0].rect.width == pytest.approx(w + 2 * ofs, abs=0.6)
    assert res.preflight["profile"] == "print" and res.preflight["bleed_pt"] == pytest.approx(3 * MM_TO_PT)
    assert any("PDF/X" in n for n in res.notes)
    assert doc[0].trimbox.width == pytest.approx(w, abs=0.6)
    assert doc.page_count == res.pages


def test_proof_profile_adds_diagnostics(tmp_path, sample_book):
    dest = tmp_path / "proof.pdf"
    export_pdf(sample_book, dest, "proof")
    assert "PROOF" in pymupdf.open(str(dest))[0].get_text()


def test_existing_destination_is_protected(tmp_path, sample_book):
    dest = tmp_path / "keep.pdf"
    dest.write_bytes(b"precious")
    with pytest.raises(ExportError, match="already exists"):
        export_pdf(sample_book, dest)
    assert dest.read_bytes() == b"precious"
    export_pdf(sample_book, dest, overwrite=True)
    assert dest.read_bytes()[:5] == b"%PDF-"


def test_interrupted_export_leaves_no_partial_file(tmp_path, sample_book):
    dest = tmp_path / "cancel.pdf"
    calls = {"n": 0}

    def cancel():
        calls["n"] += 1
        return calls["n"] > 8

    with pytest.raises(ExportError, match="cancelled"):
        export_pdf(sample_book, dest, cancel=cancel)
    assert not dest.exists() and not list(tmp_path.glob("*.tmp"))


def test_failed_overwrite_keeps_old_file(tmp_path, sample_book, monkeypatch):
    dest = tmp_path / "old.pdf"
    dest.write_bytes(b"old content")
    import app.exporters.pdf_exporter as pe

    def boom(*a, **k):
        raise RuntimeError("renderer crashed")
    monkeypatch.setattr("app.services.export_service.render_pdf", boom)
    with pytest.raises(RuntimeError):
        export_pdf(sample_book, dest, overwrite=True)
    assert dest.read_bytes() == b"old content" and not list(tmp_path.glob("*.tmp"))


def test_missing_folder_and_permission_errors(tmp_path, sample_book):
    with pytest.raises(ExportError, match="does not exist"):
        export_pdf(sample_book, tmp_path / "nope" / "x.pdf")
    ro = tmp_path / "ro"
    ro.mkdir()
    dest = ro / "x.pdf"
    if os.name == "nt":
        import subprocess
        subprocess.run(["icacls", str(ro), "/deny", f"{os.environ['USERNAME']}:(WD,AD)"],
                       capture_output=True)
    else:
        ro.chmod(stat.S_IREAD | stat.S_IEXEC)
    try:
        try:
            export_pdf(sample_book, dest)
            wrote = dest.exists()
        except ExportError as e:
            wrote = False
            assert "permission" in str(e).lower() or "cannot" in str(e).lower()
        if wrote:
            pytest.skip("could not make the destination read-only on this system")
    finally:
        if os.name == "nt":
            import subprocess
            subprocess.run(["icacls", str(ro), "/remove:d", os.environ["USERNAME"]],
                           capture_output=True)
        else:
            ro.chmod(stat.S_IRWXU)
    assert not list(ro.glob("*.tmp"))


def test_validation_blocks_export_with_missing_image(tmp_path, sample_book):
    sample_book.sections[2].blocks.append(ImageBlock(asset="ghost"))
    dest = tmp_path / "bad.pdf"
    with pytest.raises(ValidationBlocked) as ei:
        export_pdf(sample_book, dest)
    assert any(i.code == "missing_asset" for i in ei.value.report.errors)
    assert not dest.exists()
    res = export_pdf(sample_book, dest, allow_errors=True)  # user reviewed the issues
    assert res.validation.errors and dest.exists()


def test_other_formats(tmp_path, sample_book, png_bytes):
    a = sample_book.add_asset(png_bytes, "pic.png")
    sample_book.sections[2].blocks.append(ImageBlock(asset=a.id, alt="alt", caption=[Run("Fig")]))
    sample_book.sections[2].blocks.append(TableBlock(rows=[["h", "g"], ["1", "2"]]))
    r = export_other(sample_book, tmp_path / "b.html", "html")
    html = (tmp_path / "b.html").read_text(encoding="utf-8")
    assert "Cartea de test" in html and "data:image/png;base64" in html and "Ștefan" in html
    r = export_other(sample_book, tmp_path / "b.md", "markdown")
    md = (tmp_path / "b.md").read_text(encoding="utf-8")
    assert "# Cartea de test" in md and "| h | g |" in md
    assert (tmp_path / "b_images").is_dir() and any((tmp_path / "b_images").iterdir())
    r = export_other(sample_book, tmp_path / "b.docx", "docx")
    from docx import Document
    d = Document(str(tmp_path / "b.docx"))
    texts = [p.text for p in d.paragraphs]
    assert any("Ștefan" in t for t in texts) and d.core_properties.title == "Cartea de test"
    assert len(d.tables) == 1 and len(d.inline_shapes) == 1
    with pytest.raises(ExportError):
        export_other(sample_book, tmp_path / "b.html", "html")
