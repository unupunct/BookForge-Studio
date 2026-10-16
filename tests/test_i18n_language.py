import pytest
from docx import Document

from app.domain.book import BookProject
from app.domain.content_blocks import ParagraphBlock, PoetryBlock, RecipeBlock, Run, ExerciseBlock
from app.domain.sections import Section
from app.i18n import exercise_label, tr
from app.importers.base import ImportOptions
from app.layout.model import Glyphs
from app.layout.pagination import LayoutEngine
from app.services.import_service import run_import
from app.services.language_detect import detect_language
from app.services.project_service import NewBookSpec, create_from_manuscript
from app.ui.dialogs.manuscript_dialog import ManuscriptDialog
from tests.conftest import LOREM
from tests.test_ui import qapp, wait_until, window  # noqa: F401  (fixtures)

RO = ("Toată viața am spus că mai târziu mă odihnesc. Dar timpul nu așteaptă pe nimeni și noi rămânem cu "
      "ceea ce nu am apucat să facem. Poate că viața este doar o ușă care se deschide încet și nu ne "
      "întreabă dacă suntem pregătiți să intrăm în ea. ") * 3
EN = ("The old house stood at the end of the road and nobody had lived there for years. When the wind came "
      "from the sea it moved through the empty rooms as if it was looking for someone that it knew. ") * 3


def poem_book(lang, poems=3):
    b = BookProject()
    b.metadata.language = lang
    b.sections = [Section("P", "chapter", blocks=[
        PoetryBlock(title=str(i), text="\n".join(f"vers {k} din {i}" for k in range(70))) for i in range(poems)])]
    return b


def glyph_texts(res):
    return [g.text for p in res.pages for g in p.items if isinstance(g, Glyphs)]


def test_translation_table_and_fallbacks():
    assert tr("continued", "ro") == "(continuare)" and tr("continued", "ro-RO") == "(continuare)"
    assert tr("continued", "en") == "(continued)" and tr("continued", "xx") == "(continued)"
    assert tr("continued", None) == "(continued)" and tr("contents", "ro") == "Cuprins"
    assert exercise_label("definition", "ro") == "Definiție" and exercise_label("custom", "ro") == "Custom"
    for key in ("continued", "contents", "poem_index", "chapter", "part", "ingredients", "method"):
        assert all(tr(key, c) for c in ("en", "ro", "fr", "de", "es", "it", "pt", "hu", "pl", "nl"))


def test_romanian_poem_continuation_marker():
    res = LayoutEngine().build(poem_book("ro"))
    texts = glyph_texts(res)
    assert "(continuare)" in texts and "(continued)" not in texts
    en = glyph_texts(LayoutEngine().build(poem_book("en")))
    assert "(continued)" in en and "(continuare)" not in en
    fr = glyph_texts(LayoutEngine().build(poem_book("fr")))
    assert "(suite)" in fr


def test_romanian_labels_in_recipes_chapters_and_poem_index():
    b = BookProject()
    b.metadata.language = "ro"
    b.settings.chapter_number_style = "arabic"
    b.settings.chapter_label = "Chapter {n}"
    b.settings.poem_index = "titles"
    b.sections = [Section("Cuprins", "toc", in_toc=False),
                  Section("Retete", "chapter", blocks=[
                      RecipeBlock(title="Ciorbă", prep_time="10 min", cook_time="20 min", servings="4",
                                  ingredients=["apă"], instructions=["fierbe"], notes="gustoasă"),
                      ExerciseBlock(kind="note", title="x", runs=[Run("text")]),
                      PoetryBlock(title="Doina", text="a\nb")])]
    res = LayoutEngine().build(b)
    shown = " ".join(glyph_texts(res))
    for word in ("Capitolul 1", "Pregătire: 10 min", "Gătire: 20 min", "Porții: 4", "Ingrediente", "Mod de preparare",
                 "Note: ", "gustoasă", "Notă: x", "Index de poezii"):
        assert word in shown, word
    assert "Ingredients" not in shown and "Chapter 1" not in shown
    b.metadata.language = "en"
    shown_en = " ".join(glyph_texts(LayoutEngine().build(b)))
    assert "Chapter 1" in shown_en and "Prep: 10 min" in shown_en and "Index of poems" in shown_en


def test_language_detection():
    assert detect_language(RO) == "ro" and detect_language(EN) == "en"
    assert detect_language("prea scurt") is None and detect_language("") is None
    assert detect_language("1234 5678 " * 100) is None
    fr = "Le chat est sur la table et il regarde les oiseaux dans le jardin avec une grande attention pour la journée. " * 4
    assert detect_language(fr) == "fr"


def test_import_detects_language_and_book_follows(tmp_path):
    f = tmp_path / "ro.docx"
    d = Document()
    d.add_heading("Capitolul 1", 1)
    d.add_paragraph(RO)
    d.save(f)
    staged = run_import(f, ImportOptions())
    assert staged.result.language == "ro"
    assert any("detected as 'ro'" in i.message for i in staged.result.report.issues)
    book = create_from_manuscript(NewBookSpec(profile="poetry", template="poetry_ballad_epic"), staged, ImportOptions())
    assert book.metadata.language == "ro"
    explicit = create_from_manuscript(NewBookSpec(profile="novel", language="en"), run_import(f), ImportOptions())
    assert explicit.metadata.language == "en"  # an explicit choice wins over detection
    toc = next((s for s in book.sections if s.kind == "toc"), None)
    assert toc is None or toc.title == "Cuprins"


def test_collection_import_gets_romanian_contents_title(tmp_path):
    from tests.test_import_collection import collection_docx
    f = tmp_path / "c.docx"
    collection_docx(f, n=12)
    book = create_from_manuscript(NewBookSpec(profile="poetry", template="poetry_ballad_epic"),
                                  run_import(f, ImportOptions()), ImportOptions())
    assert book.metadata.language == "ro"  # the fixture poems are Romanian
    assert next(s for s in book.sections if s.kind == "toc").title == "Cuprins"


def test_manuscript_dialog_prefills_detected_language(window, tmp_path):
    f = tmp_path / "ro.docx"
    d = Document()
    d.add_heading("Capitolul 1", 1)
    d.add_paragraph(RO)
    d.save(f)
    dlg = ManuscriptDialog(window, window.state.templates, "", str(f))
    assert wait_until(lambda: dlg.staged is not None)
    assert dlg.language.currentData() == "ro" and dlg.spec().language == "ro"
    dlg.language.setCurrentIndex(dlg.language.findData("fr"))
    assert dlg.spec().language == "fr"
    dlg.close()
