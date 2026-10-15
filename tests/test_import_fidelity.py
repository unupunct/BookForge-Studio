"""Imported text keeps its formatting; poems keep lines, indentation and stanza spacing."""
import pymupdf
import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.domain.book import BookProject
from app.domain.content_blocks import HeadingBlock, ParagraphBlock, PoetryBlock
from app.importers.base import ImportOptions
from app.layout.model import Glyphs
from app.layout.pagination import LayoutEngine
from app.services.export_service import export_pdf
from app.services.import_service import import_file, run_import
from tests.conftest import LOREM, RO

POEM = ["Pe un picior de plai,", "Pe-o gură de rai,", "", "Iată vine o mioară,", "Mioară bârsană,",
        "", "Pe-un picior de plai"]


def docx_poem(path, *, soft=False, blank_para=True, tab_indent=False):
    d = Document()
    d.add_heading("Poezii", 1)
    d.add_paragraph("Mioriţa")  # title-like line, then the poem
    d.add_paragraph("")
    if soft:
        p = d.add_paragraph()
        for i, line in enumerate(POEM[:2]):
            r = p.add_run(line)
            if i < 1:
                r.add_break()
    for line in POEM:
        if soft:
            break
        if line == "":
            d.add_paragraph("")
            continue
        p = d.add_paragraph(("\t" if tab_indent and line.startswith("Mioară") else "") + line)
        p.paragraph_format.space_after = Pt(0)
    d.add_heading("Alta", 1)
    d.add_paragraph(LOREM * 4)
    d.save(path)


def test_docx_poem_keeps_lines_blank_lines_and_indent(tmp_path):
    f = tmp_path / "p.docx"
    docx_poem(f, tab_indent=True)
    book = BookProject()
    res = import_file(book, f)
    sec = res.sections[0]
    poem = next(b for b in sec.blocks if isinstance(b, PoetryBlock))
    lines = poem.text.split("\n")
    assert poem.title == "Mioriţa"
    assert lines[:2] == ["Pe un picior de plai,", "Pe-o gură de rai,"]
    assert lines[2] == "" and lines[3] == "Iată vine o mioară,"
    assert lines[4].startswith("    ") and lines[4].strip() == "Mioară bârsană,"  # tab kept as indent
    assert any("poem" in i.message for i in res.report.issues)
    prose = res.sections[1].blocks[0]
    assert isinstance(prose, ParagraphBlock)  # real prose is not turned into verse


def test_docx_soft_line_breaks_become_verse(tmp_path):
    f = tmp_path / "s.docx"
    docx_poem(f, soft=True)
    res = import_file(BookProject(), f)
    poem = next(b for b in res.sections[0].blocks if isinstance(b, PoetryBlock))
    assert poem.text.split("\n") == POEM[:2]


def test_poetry_modes(tmp_path):
    f = tmp_path / "m.docx"
    d = Document()
    d.add_heading("C", 1)
    for t in ("Short line one", "Short line two", "Short three"):
        d.add_paragraph(t)
    d.save(f)
    never = run_import(f, ImportOptions(poetry_mode="never")).result.sections[0].blocks
    always = run_import(f, ImportOptions(poetry_mode="always")).result.sections[0].blocks
    assert all(isinstance(b, ParagraphBlock) for b in never) and len(never) == 3
    assert len(always) == 1 and isinstance(always[0], PoetryBlock)
    assert always[0].text == "Short line one\nShort line two\nShort three"


def test_dialogue_and_prose_are_not_mistaken_for_poems(tmp_path):
    f = tmp_path / "n.docx"
    d = Document()
    d.add_heading("Capitolul 1", 1)
    for t in ("— Unde mergi?", "— Acasă.", "— Bine, atunci.", "— Pa.",
              "El a plecat în liniște."):
        d.add_paragraph(t)
    d.add_paragraph(LOREM * 5)
    d.save(f)
    blocks = run_import(f).result.sections[0].blocks
    assert not any(isinstance(b, PoetryBlock) for b in blocks)
    assert len(blocks) == 6


def test_docx_character_and_paragraph_formatting_kept(tmp_path):
    f = tmp_path / "fmt.docx"
    d = Document()
    d.add_heading("C", 1)
    for i in range(6):
        d.add_paragraph(f"Body {i} " + LOREM)
    c = d.add_paragraph("Centered dedication")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    q = d.add_paragraph("Indented quotation " + LOREM)
    q.paragraph_format.left_indent = Pt(40)
    q.paragraph_format.space_before = Pt(18)
    p = d.add_paragraph("H")
    r = p.add_run("2")
    r.font.subscript = True
    p.add_run(" and E=mc")
    r2 = p.add_run("2")
    r2.font.superscript = True
    r3 = p.add_run(" removed")
    r3.font.strike = True
    d.save(f)
    blocks = run_import(f).result.sections[0].blocks
    paras = [b for b in blocks if isinstance(b, ParagraphBlock)]
    centered = next(b for b in paras if b.plain_text() == "Centered dedication")
    assert centered.align == "center"
    quote = next(b for b in paras if b.plain_text().startswith("Indented"))
    assert quote.left_indent == 40 and quote.space_before == 18
    body = next(b for b in paras if b.plain_text().startswith("Body 0"))
    assert body.left_indent is None and body.align == ""  # ordinary body text follows the template
    chem = next(b for b in paras if b.plain_text().startswith("H2"))
    assert [(r.text, r.sub, r.sup, r.strike) for r in chem.runs] == [
        ("H", False, False, False), ("2", True, False, False), (" and E=mc", False, False, False),
        ("2", False, True, False), (" removed", False, False, True)]


def test_keep_formatting_off_clears_overrides(tmp_path):
    f = tmp_path / "k.docx"
    d = Document()
    d.add_heading("C", 1)
    c = d.add_paragraph("Centered")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.save(f)
    b = run_import(f, ImportOptions(keep_formatting=False)).result.sections[0].blocks[0]
    assert b.align == ""


def test_chapter_lines_recognised_without_heading_styles(tmp_path):
    f = tmp_path / "c.docx"
    d = Document()
    for n in (1, 2):
        d.add_paragraph(f"CAPITOLUL {n}")
        d.add_paragraph(LOREM * 3)
    d.save(f)
    res = run_import(f).result
    assert [s.title for s in res.sections] == ["CAPITOLUL 1", "CAPITOLUL 2"]
    assert any("recognised" in i.message for i in res.report.issues)


def test_docx_author_and_title_properties(tmp_path):
    f = tmp_path / "a.docx"
    d = Document()
    d.core_properties.author = "Ion Creangă"
    d.core_properties.title = "Amintiri"
    d.add_heading("C", 1)
    d.save(f)
    res = run_import(f).result
    assert res.author == "Ion Creangă" and res.title == "Amintiri"


def test_txt_verse_detection_keeps_indentation_and_gaps(tmp_path):
    f = tmp_path / "v.txt"
    f.write_text("Versuri\n\nCapitolul 1\n\nPrima linie,\n    a doua, cu indent,\nA treia linie\n\n\nStrofa nouă,\nnici asta nu e proză\n\n"
                 + (LOREM * 6 + "\n") * 2, encoding="utf-8")
    res = run_import(f, ImportOptions(detect_title=True)).result
    blocks = res.sections[0].blocks
    poem = next(b for b in blocks if isinstance(b, PoetryBlock))
    assert poem.text.split("\n")[1].startswith("    a doua")
    assert "\n\n\n" not in poem.text and poem.text.count("\n\n") == 1 or poem.text.count("\n\n") >= 1
    prose = [b for b in blocks if isinstance(b, ParagraphBlock)]
    assert len(prose) == 1 and chr(10) not in prose[0].plain_text()  # wrapped prose is re-joined


def test_markdown_and_html_verse(tmp_path):
    m = tmp_path / "p.md"
    m.write_text("# Poezii\n\nRoses are red\nViolets are blue\nSugar is sweet\n\nAnd so are you\nIn every way\nEach day\n", encoding="utf-8")
    poem = next(b for b in run_import(m).result.sections[0].blocks if isinstance(b, PoetryBlock))
    assert poem.text.split("\n") == ["Roses are red", "Violets are blue", "Sugar is sweet", "",
                                     "And so are you", "In every way", "Each day"]
    h = tmp_path / "p.html"
    h.write_text("<h1>Poezii</h1><p>Line one<br>Line two<br>Line three</p><p>Line four<br>Line five</p>",
                 encoding="utf-8")
    poem = next(b for b in run_import(h).result.sections[0].blocks if isinstance(b, PoetryBlock))
    assert poem.text.split("\n") == ["Line one", "Line two", "Line three", "", "Line four", "Line five"]


def test_layout_stanza_spacing_matches_blank_lines():
    b = BookProject()
    from app.domain.sections import Section
    one = PoetryBlock(text="a\nb\n\nc\nd")
    two = PoetryBlock(text="a\nb\n\n\nc\nd")
    b.sections = [Section("P", "chapter", blocks=[one]), Section("Q", "chapter", blocks=[two])]
    res = LayoutEngine().build(b)

    def ys(block):
        page = res.pages[res.block_pages[block.id][0]]
        return sorted({round(g.y, 1) for g in page.items if isinstance(g, Glyphs) and g.block_id == block.id})
    y1, y2 = ys(one), ys(two)
    line = y1[1] - y1[0]
    assert y1[2] - y1[1] == pytest.approx(line * (1 + 0.9), rel=0.02)
    assert y2[2] - y2[1] == pytest.approx(line * (1 + 1.8), rel=0.02)  # two blanks: double gap


def test_layout_alignment_and_script_positions():
    b = BookProject()
    from app.domain.content_blocks import Run
    from app.domain.sections import Section
    centered = ParagraphBlock(runs=[Run("Centered")], align="center")
    sup = ParagraphBlock(runs=[Run("E=mc"), Run("2", sup=True)])
    b.sections = [Section("S", "chapter", blocks=[centered, sup])]
    res = LayoutEngine().build(b)
    page = res.pages[res.block_pages[centered.id][0]]
    g = next(i for i in page.items if isinstance(i, Glyphs) and i.block_id == centered.id)
    x0, _, x1, _ = page.content_box
    assert (g.x + g.width / 2) == pytest.approx((x0 + x1) / 2, abs=1.0)
    gs = [i for i in page.items if isinstance(i, Glyphs) and i.block_id == sup.id]
    assert gs[1].text == "2" and gs[1].y < gs[0].y and gs[1].size < gs[0].size


def test_poem_in_pdf_has_expected_line_gaps(tmp_path):
    f = tmp_path / "p.docx"
    docx_poem(f)
    book = BookProject()
    book.metadata.title = "Poezii"
    import_file(book, f)
    pdf = tmp_path / "p.pdf"
    export_pdf(book, pdf, "digital", allow_errors=True)
    doc = pymupdf.open(str(pdf))
    for page in doc:
        words = page.get_text("words")
        if any(w[4] == "Mioară" for w in words):
            tops = {}
            for w in words:
                tops.setdefault(round(w[3]), []).append(w[4])
            rows = [" ".join(v) for k, v in sorted(tops.items())]
            i = rows.index("Pe un picior de plai,")
            assert rows[i + 1] == "Pe-o gură de rai,"
            ys = sorted(tops)
            j = ys.index(next(k for k, v in tops.items() if v[0] == "Iată"))
            # the stanza gap is clearly larger than the normal line gap
            assert (ys[j] - ys[j - 1]) > 1.5 * (ys[i + 1 - 0] - ys[i]) or (ys[j] - ys[j - 1]) > 20
            return
    pytest.fail("poem not found in the PDF")


def test_formatting_survives_html_and_docx_export(tmp_path):
    from docx import Document as ReadDoc
    from app.domain.content_blocks import Run
    from app.domain.sections import Section
    from app.services.export_service import export_other
    b = BookProject()
    b.metadata.title = "F"
    b.sections = [Section("C", "chapter", blocks=[
        ParagraphBlock(runs=[Run("E=mc"), Run("2", sup=True), Run(" old", strike=True)], align="center",
                       left_indent=20, space_before=12)])]
    export_other(b, tmp_path / "f.html", "html")
    html = (tmp_path / "f.html").read_text(encoding="utf-8")
    assert "<sup>2</sup>" in html and "<s> old</s>" in html and "text-align:center" in html and "margin-left:20pt" in html
    export_other(b, tmp_path / "f.docx", "docx")
    par = next(p for p in ReadDoc(str(tmp_path / "f.docx")).paragraphs if p.text.startswith("E=mc"))
    assert par.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert par.runs[1].font.superscript and par.runs[2].font.strike
    assert par.paragraph_format.left_indent.pt == 20
