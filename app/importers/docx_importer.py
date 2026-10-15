"""DOCX import with python-docx plus direct OOXML access for what it does not expose."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree

from ..domain.content_blocks import (Block, CodeBlock, FootnoteBlock, HeadingBlock, ImageBlock,
                                     ListBlock, PageBreakBlock, ParagraphBlock, QuoteBlock, Run,
                                     TableBlock, merge_runs, runs_text)
from ..domain.sections import Section
from .structure import apply_header_policy, extract_title_block
from .postprocess import (Blank, convert_poetry, normalise_paragraph_format,
                          promote_chapter_headings)
from .base import (ImageSink, ImportFailure, ImportOptions, ImportReport, ImportResult,
                   split_into_sections, stem_title)

MONO_FONTS = {"consolas", "courier new", "courier", "lucida console", "monaco", "menlo",
              "source code pro", "cascadia code", "cascadia mono"}
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
MAX_UNCOMPRESSED = 1024 * 1024 * 1024


class _Converter:
    def __init__(self, doc: Any, sink: ImageSink, report: ImportReport):
        self.doc, self.sink, self.report = doc, sink, report
        self.blocks: list[Block] = []
        self.title: str | None = None
        self.footnotes: list[tuple[str, str]] = []  # (label, text)
        self.fn_text = self._note_texts("footnotes")
        self.en_text = self._note_texts("endnotes")
        self.numbering = self._numbering_formats()
        self.pending_caption: list[Run] | None = None

    # -- helpers -------------------------------------------------------------
    def _note_texts(self, kind: str) -> dict[str, str]:
        out: dict[str, str] = {}
        try:
            for rel in self.doc.part.rels.values():
                if rel.reltype.endswith("/" + kind) and not rel.is_external:
                    root = etree.fromstring(rel.target_part.blob)
                    for note in root.findall(qn(f"w:{kind[:-1]}")):
                        nid = note.get(qn("w:id"))
                        if nid is None or int(nid) <= 0:
                            continue
                        out[nid] = " ".join("".join(t.text or "" for t in note.iter(qn("w:t"))).split())
        except Exception as e:
            self.report.add("warning", f"{kind.title()} could not be read: {e}")
        return out

    def _numbering_formats(self) -> dict[tuple[str, str], str]:
        fmts: dict[tuple[str, str], str] = {}
        try:
            root = self.doc.part.numbering_part.element
        except Exception:
            return fmts
        abstract: dict[str, dict[str, str]] = {}
        for an in root.findall(qn("w:abstractNum")):
            aid = an.get(qn("w:abstractNumId"))
            for lvl in an.findall(qn("w:lvl")):
                nf = lvl.find(qn("w:numFmt"))
                abstract.setdefault(aid, {})[lvl.get(qn("w:ilvl"))] = nf.get(qn("w:val")) if nf is not None else "bullet"
        for num in root.findall(qn("w:num")):
            nid = num.get(qn("w:numId"))
            ref = num.find(qn("w:abstractNumId"))
            if ref is not None:
                for ilvl, fmt in abstract.get(ref.get(qn("w:val")), {}).items():
                    fmts[(nid, ilvl)] = fmt
        return fmts

    def _list_info(self, p: Paragraph) -> tuple[bool, bool]:
        """(is_list_item, ordered)."""
        style = (p.style.name if p.style is not None else "") or ""
        ppr = p._p.pPr
        num = ppr.numPr if ppr is not None else None
        if num is not None and num.numId is not None and int(num.numId.val) != 0:
            ilvl = str(num.ilvl.val) if num.ilvl is not None else "0"
            fmt = self.numbering.get((str(num.numId.val), ilvl))
            return True, fmt not in (None, "bullet", "none")
        if style.lower().startswith("list number"):
            return True, True
        if style.lower().startswith("list bullet"):
            return True, False
        return False, False

    # -- runs ------------------------------------------------------------------
    def _run_format(self, r: Any) -> dict[str, Any]:
        rpr = r.find(qn("w:rPr"))
        fmt: dict[str, Any] = {"bold": False, "italic": False, "underline": False, "code": False}
        if rpr is not None:
            def on(tag: str) -> bool:
                el = rpr.find(qn(tag))
                if el is None:
                    return False
                v = el.get(qn("w:val"))
                return v not in ("0", "false", "off", "none")
            fmt["bold"], fmt["italic"], fmt["underline"] = on("w:b"), on("w:i"), on("w:u")
            fmt["strike"] = on("w:strike") or on("w:dstrike")
            va = rpr.find(qn("w:vertAlign"))
            va = va.get(qn("w:val")) if va is not None else None
            fmt["sup"], fmt["sub"] = va == "superscript", va == "subscript"
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None:
                name = (fonts.get(qn("w:ascii")) or fonts.get(qn("w:hAnsi")) or "").lower()
                fmt["code"] = name in MONO_FONTS
        return fmt

    def _run_text(self, r: Any) -> str:
        out = []
        for ch in r:
            tag = etree.QName(ch).localname
            if tag == "t":
                out.append(ch.text or "")
            elif tag == "tab":
                out.append("    ")
            elif tag == "br" and ch.get(qn("w:type")) in (None, "textWrapping"):
                out.append("\n")
            elif tag == "noBreakHyphen":
                out.append("-")
            elif tag == "sym":
                self.report.add("warning", "Symbol characters (w:sym) were not imported")
        return "".join(out)

    def _paragraph_content(self, p: Paragraph) -> tuple[list[list[Run] | str], list[ImageBlock]]:
        """Split paragraph content at page breaks. Items are run lists or the string 'PAGEBREAK'."""
        segments: list[list[Run] | str] = [[]]
        images: list[ImageBlock] = []

        def add_run(r: Any, link: str | None) -> None:
            for br in r.findall(qn("w:br")):
                if br.get(qn("w:type")) == "page":
                    segments.append("PAGEBREAK")
                    segments.append([])
            for ref in r.findall(qn("w:footnoteReference")):
                nid = ref.get(qn("w:id"))
                if nid in self.fn_text:
                    label = f"[{len(self.footnotes) + 1}]"
                    self.footnotes.append((label, self.fn_text[nid]))
                    segments[-1].append(Run(label, link=None))  # type: ignore[union-attr]
            for ref in r.findall(qn("w:endnoteReference")):
                nid = ref.get(qn("w:id"))
                if nid in self.en_text:
                    label = f"[{len(self.footnotes) + 1}]"
                    self.footnotes.append((label, self.en_text[nid]))
                    segments[-1].append(Run(label))  # type: ignore[union-attr]
            for d in r.iter(qn("w:drawing")):
                self._drawing(d, images)
            if r.find(qn("w:pict")) is not None or r.find(qn("w:object")) is not None:
                self.report.add("warning", "Legacy VML pictures/embedded objects were not imported")
            text = self._run_text(r)
            if text:
                f = self._run_format(r)
                segments[-1].append(Run(text, f["bold"], f["italic"], f["underline"], f["code"], link,
                                        f.get("sup", False), f.get("sub", False), f.get("strike", False)))  # type: ignore[union-attr]

        for child in p._p.iterchildren():
            tag = etree.QName(child).localname
            if tag == "r":
                add_run(child, None)
            elif tag == "hyperlink":
                rid = child.get(qn("r:id"))
                link = None
                if rid and rid in self.doc.part.rels:
                    link = self.doc.part.rels[rid].target_ref
                elif child.get(qn("w:anchor")):
                    self.report.add("info", "Internal bookmark hyperlinks were imported as plain text")
                for r in child.findall(qn("w:r")):
                    add_run(r, link)
            elif tag in ("ins", "smartTag", "sdt"):
                if tag == "ins":
                    self.report.add("info", "Tracked insertions were accepted into the text")
                for r in child.iter(qn("w:r")):
                    add_run(r, None)
            elif tag == "del":
                self.report.add("warning", "Tracked deletions were left out of the imported text")
            elif tag in ("oMath", "oMathPara"):
                self.report.add("warning", "Equations (OMML) were not imported")
        return segments, images

    def _drawing(self, d: Any, images: list[ImageBlock]) -> None:
        blips = d.findall(".//{%s}blip" % A_NS)
        if not blips:
            self.report.add("warning", "A drawing/shape without an embedded picture was skipped")
            return
        doc_pr = d.find(".//{%s}docPr" % WP_NS)
        alt = (doc_pr.get("descr") or doc_pr.get("title") or "") if doc_pr is not None else ""
        for blip in blips:
            rid = blip.get("{%s}embed" % R_NS)
            if not rid or rid not in self.doc.part.related_parts:
                self.report.add("warning", "An image reference could not be resolved")
                continue
            part = self.doc.part.related_parts[rid]
            name = Path(str(part.partname)).name
            asset = self.sink(part.blob, name)
            if asset is None:
                self.report.add("warning", f"Image '{name}' has an unsupported format and was not imported")
                continue
            images.append(ImageBlock(asset=asset, alt=alt))
            self.report.count("images")

    # -- blocks ----------------------------------------------------------------
    def convert(self) -> None:
        body = self.doc.element.body
        self._children(body)
        self._flush_caption()

    def _children(self, container: Any) -> None:
        for child in container.iterchildren():
            tag = etree.QName(child).localname
            if tag == "p":
                self._paragraph(Paragraph(child, self.doc))
            elif tag == "tbl":
                self._table(Table(child, self.doc))
            elif tag == "sdt":
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    self._children(content)
            elif tag == "sectPr":
                continue
            elif tag in ("bookmarkStart", "bookmarkEnd", "proofErr"):
                continue
            else:
                self.report.add("info", f"Unsupported body element <{tag}> skipped")

    @staticmethod
    def _direct_format(p: Paragraph) -> dict:
        """Alignment, indents and spacing set on the paragraph (or its style for alignment)."""
        pf = p.paragraph_format
        out: dict = {}
        al = pf.alignment
        if al is None and p.style is not None:
            al = p.style.paragraph_format.alignment
        name = {0: "left", 1: "center", 2: "right", 3: "justify"}.get(int(al) if al is not None else -1, "")
        if name:
            out["align"] = name
        for attr, key in (("left_indent", "left_indent"), ("first_line_indent", "first_indent"),
                          ("space_before", "space_before"), ("space_after", "space_after")):
            v = getattr(pf, attr)
            if v is not None:
                out[key] = round(v.pt, 1)
        return out

    def _flush_caption(self) -> None:
        if self.pending_caption:
            self.blocks.append(ParagraphBlock(runs=self.pending_caption, style="caption"))
        self.pending_caption = None

    def _paragraph(self, p: Paragraph) -> None:
        style = (p.style.name if p.style is not None else "") or ""
        if style.lower() not in ("caption", "figure caption", "table caption"):
            self._flush_caption()
        ppr = p._p.pPr
        if ppr is not None and ppr.find(qn("w:pageBreakBefore")) is not None:
            self.blocks.append(PageBreakBlock())
        if p._p.xpath(".//w:txbxContent"):
            self.report.add("warning", "Text boxes were not imported")
        if p._p.xpath(".//w:fldSimple|.//w:instrText"):
            self.report.add("info", "Fields (page numbers, TOC, cross-references) were imported as their cached text only")
        segments, images = self._paragraph_content(p)
        is_list, ordered = self._list_info(p)
        low = style.lower()
        fmt = self._direct_format(p)
        if not images and all(sg == [] for sg in segments):
            self.blocks.append(Blank())  # an empty paragraph: spacing the author chose
        for seg in segments:
            if seg == "PAGEBREAK":
                self.blocks.append(PageBreakBlock())
                self.report.count("page breaks")
                continue
            runs = merge_runs(seg)  # type: ignore[arg-type]
            if not runs and not images:
                continue
            text = runs_text(runs)
            if runs and not text.strip():
                runs = []
            if not runs:
                continue
            m = re.match(r"heading\s*(\d)", low)
            if m:
                self.blocks.append(HeadingBlock(level=min(int(m.group(1)), 6), runs=runs))
                self.report.count("headings")
            elif low == "title":
                if self.title is None:
                    self.title = text.strip()
                    self.report.add("info", f"'Title' paragraph used as document title: '{self.title}'")
                else:
                    self.blocks.append(HeadingBlock(level=1, runs=runs))
            elif low == "subtitle":
                self.blocks.append(ParagraphBlock(runs=runs, style="subtitle"))
            elif low in ("caption", "figure caption", "table caption"):
                last = self.blocks[-1] if self.blocks else None
                if isinstance(last, ImageBlock) and not last.caption:
                    last.caption = runs
                else:
                    self.pending_caption = runs
                self.report.count("captions")
            elif low in ("quote", "intense quote", "block text", "block quote"):
                self.blocks.append(QuoteBlock(runs=runs))
            elif is_list:
                last = self.blocks[-1] if self.blocks else None
                if isinstance(last, ListBlock) and last.ordered == ordered:
                    last.items.append(runs)
                else:
                    self.blocks.append(ListBlock(ordered=ordered, items=[runs]))
                self.report.count("list items")
            elif low in ("code", "source code", "html code", "preformatted text") or \
                    (all(r.code for r in runs)):
                last = self.blocks[-1] if self.blocks else None
                if isinstance(last, CodeBlock):
                    last.text += "\n" + text
                else:
                    self.blocks.append(CodeBlock(text=text))
            else:
                self.blocks.append(ParagraphBlock(runs=runs, **fmt))
                self.report.count("paragraphs")
        for img in images:
            self.blocks.append(img)
        if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
            self.blocks.append(PageBreakBlock())
            self.report.add("info", "Section breaks were converted to page breaks")

    def _table(self, t: Table) -> None:
        rows: list[list[str]] = []
        merged = False
        nested = False
        for tr in t._tbl.iterchildren(qn("w:tr")):
            row: list[str] = []
            for tc in tr.iterchildren(qn("w:tc")):
                span = 1
                tcpr = tc.find(qn("w:tcPr"))
                if tcpr is not None:
                    gs = tcpr.find(qn("w:gridSpan"))
                    if gs is not None:
                        span = int(gs.get(qn("w:val")) or 1)
                    if tcpr.find(qn("w:vMerge")) is not None:
                        merged = True
                if tc.find(qn("w:tbl")) is not None:
                    nested = True
                text = "\n".join("".join(x.text or "" for x in p.iter(qn("w:t")))
                                 for p in tc.iterchildren(qn("w:p"))).strip()
                row.append(text)
                if span > 1:
                    merged = True
                    row.extend([""] * (span - 1))
            rows.append(row)
        if merged:
            self.report.add("warning", "Merged table cells were flattened into separate cells")
        if nested:
            self.report.add("warning", "Nested tables were flattened to the text of their parent cell")
        if not rows:
            return
        cap = self.pending_caption or []
        self.pending_caption = None
        self.blocks.append(TableBlock(rows=rows, header_rows=1, caption=cap))
        self.report.count("tables")


def import_docx(path: Path, raw: bytes, options: ImportOptions, sink: ImageSink) -> ImportResult:
    report = ImportReport(source=Path(path).name, format="DOCX")
    try:
        with zipfile.ZipFile(Path(path)) as z:
            total = sum(i.file_size for i in z.infolist())
            if total > MAX_UNCOMPRESSED:
                raise ImportFailure("The DOCX file expands to an unreasonable size and was rejected")
            names = set(z.namelist())
            if "word/vbaProject.bin" in names:
                report.add("warning", "Embedded macros were ignored (never executed)")
            if "word/comments.xml" in names:
                report.add("warning", "Comments were not imported")
        doc = Document(str(path))
    except ImportFailure:
        raise
    except Exception as e:
        raise ImportFailure(f"The DOCX file could not be read: {e}") from e
    conv = _Converter(doc, sink, report)
    try:
        conv.convert()
    except Exception as e:
        raise ImportFailure(f"The DOCX content could not be converted: {e}") from e
    blocks = conv.blocks
    while blocks and isinstance(blocks[-1], (PageBreakBlock, Blank)):
        blocks.pop()
    GENERIC_AUTHORS = {"python-docx", "unknown", "user", "admin", "author", "microsoft office user", "office",
                       "owner", "default"}
    subtitle = description = None
    if conv.title is None:  # no Title style: centred opening lines are the title page text
        t, subtitle, description, blocks = extract_title_block(blocks, report)
        conv.title = t or conv.title
    while blocks and isinstance(blocks[0], (PageBreakBlock, Blank)):  # leading breaks only leave blank pages
        blocks.pop(0)
    blocks = promote_chapter_headings(blocks, options.split_level, options.chapter_pattern, report)
    mode = "always" if options.preserve_line_breaks else options.poetry_mode
    blocks = convert_poetry(blocks, mode, report, signature=options.signature_lines)
    apply_header_policy(blocks, options.header_lines, report)
    if options.keep_formatting:
        normalise_paragraph_format(blocks)
    else:
        for b in blocks:
            if isinstance(b, ParagraphBlock):
                b.align, b.left_indent, b.first_indent, b.space_before, b.space_after = "", None, None, None, None
    sections = split_into_sections(blocks, options.split_level, conv.title or stem_title(path), report)
    if conv.footnotes:
        notes = Section("Notes", "back_matter", blocks=[
            FootnoteBlock(label=lbl, runs=[Run(txt)]) for lbl, txt in conv.footnotes])
        sections.append(notes)
        report.add("info", f"{len(conv.footnotes)} footnotes/endnotes were collected into a 'Notes' section at the end")
        report.count("notes", len(conv.footnotes))
    props = doc.core_properties
    title = conv.title or (props.title or None)
    author = props.author if props.author and props.author.strip().lower() not in GENERIC_AUTHORS else None
    res = ImportResult(sections, report, title, author)
    res.subtitle, res.description = subtitle, description
    return res
