"""Editable DOCX export of the manuscript structure."""
from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Mm, Pt

from ..domain.book import BookProject
from ..domain.content_blocks import (Block, CodeBlock, ExerciseBlock, FootnoteBlock, HeadingBlock,
                                     ImageBlock, ListBlock, PageBreakBlock, ParagraphBlock,
                                     PoetryBlock, QuoteBlock, RecipeBlock, Run, SceneBreakBlock,
                                     TableBlock, runs_text)
from ..domain.sections import GENERATED_KINDS, Section
from ..layout.page_builder import EXERCISE_LABELS


def _add_runs(par, runs: list[Run]) -> None:
    for r in runs:
        parts = r.text.split("\n")
        for i, part in enumerate(parts):
            run = par.add_run(part)
            run.bold, run.italic, run.underline = r.bold or None, r.italic or None, r.underline or None
            if r.strike:
                run.font.strike = True
            if r.sup:  # superscript and subscript share one XML element: set only the one in use
                run.font.superscript = True
            elif r.sub:
                run.font.subscript = True
            if r.code:
                run.font.name = "Consolas"
            if i < len(parts) - 1:
                run.add_break()


_ALIGN = {"left": 0, "center": 1, "right": 2, "justify": 3}


def export_docx(book: BookProject, path: Path) -> list[str]:
    notes: list[str] = []
    doc = Document()
    md = book.metadata
    cp = doc.core_properties
    cp.title, cp.author, cp.subject, cp.language = md.title, md.author, md.subtitle, md.language
    sec = doc.sections[0]
    w, h = book.settings.page_size_pt()
    sec.page_width, sec.page_height = Pt(w), Pt(h)
    m = book.settings.margins
    sec.top_margin, sec.bottom_margin = Mm(m.top), Mm(m.bottom)
    sec.left_margin, sec.right_margin = Mm(m.inner), Mm(m.outer)
    text_w = Pt(w) - sec.left_margin - sec.right_margin

    doc.add_heading(md.title, 0)
    if md.subtitle:
        doc.add_paragraph(md.subtitle, style="Subtitle")
    if md.author:
        doc.add_paragraph(md.author)

    def block(b: Block) -> None:
        if isinstance(b, ParagraphBlock):
            par = doc.add_paragraph()
            _add_runs(par, b.runs)
            if b.align in _ALIGN:
                par.alignment = _ALIGN[b.align]
            if b.left_indent is not None:
                par.paragraph_format.left_indent = Pt(b.left_indent)
            if b.first_indent is not None:
                par.paragraph_format.first_line_indent = Pt(b.first_indent)
            if b.space_before is not None:
                par.paragraph_format.space_before = Pt(b.space_before)
            if b.space_after is not None:
                par.paragraph_format.space_after = Pt(b.space_after)
        elif isinstance(b, HeadingBlock):
            doc.add_heading(runs_text(b.runs), level=min(b.level + 1, 9))
        elif isinstance(b, ListBlock):
            style = "List Number" if b.ordered else "List Bullet"
            for it in b.items:
                _add_runs(doc.add_paragraph(style=style), it)
        elif isinstance(b, QuoteBlock):
            p = doc.add_paragraph(style="Quote")
            _add_runs(p, b.runs)
            if b.attribution:
                p.add_run("\n— " + runs_text(b.attribution))
        elif isinstance(b, CodeBlock):
            p = doc.add_paragraph()
            r = p.add_run(b.text)
            r.font.name, r.font.size = "Consolas", Pt(9)
        elif isinstance(b, PoetryBlock):
            if b.title:
                doc.add_paragraph().add_run(b.title).bold = True
            p = doc.add_paragraph()
            for i, line in enumerate(b.text.split("\n")):
                run = p.add_run(line)
                if i < len(b.text.split("\n")) - 1:
                    run.add_break()
        elif isinstance(b, ImageBlock):
            a = book.assets.get(b.asset)
            if a is None or not a.data:
                doc.add_paragraph(f"[Missing image: {b.alt}]")
                notes.append("A missing image was replaced by a note in the DOCX")
                return
            from PIL import Image
            data = a.data
            with Image.open(io.BytesIO(data)) as im:
                if im.format not in ("PNG", "JPEG", "GIF", "BMP", "TIFF"):
                    buf = io.BytesIO()
                    im.convert("RGBA").save(buf, "PNG")
                    data = buf.getvalue()
            doc.add_picture(io.BytesIO(data), width=int(text_w * min(b.width_pct, 100) / 100))
            if b.caption:
                _add_runs(doc.add_paragraph(style="Caption"), b.caption)
        elif isinstance(b, TableBlock):
            if not b.rows:
                return
            n = max(len(r) for r in b.rows)
            t = doc.add_table(rows=len(b.rows), cols=n)
            t.style = "Table Grid"
            for ri, row in enumerate(b.rows):
                for ci in range(n):
                    cell = t.cell(ri, ci)
                    cell.text = row[ci] if ci < len(row) else ""
                    if ri < b.header_rows:
                        for r in cell.paragraphs[0].runs:
                            r.bold = True
            if b.caption:
                _add_runs(doc.add_paragraph(style="Caption"), b.caption)
        elif isinstance(b, SceneBreakBlock):
            doc.add_paragraph(b.glyph).alignment = 1
        elif isinstance(b, PageBreakBlock):
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif isinstance(b, RecipeBlock):
            doc.add_heading(b.title, level=3)
            meta = [f"{k}: {v}" for k, v in (("Prep", b.prep_time), ("Cook", b.cook_time),
                                              ("Serves", b.servings)) if v]
            if meta:
                doc.add_paragraph(" · ".join(meta)).runs[0].italic = True
            if b.description:
                doc.add_paragraph(b.description)
            for i in b.ingredients:
                doc.add_paragraph(i, style="List Bullet")
            for i in b.instructions:
                doc.add_paragraph(i, style="List Number")
            if b.notes:
                doc.add_paragraph("Notes: " + b.notes)
            if b.nutrition:
                doc.add_paragraph("Nutrition (as supplied): " + b.nutrition)
        elif isinstance(b, ExerciseBlock):
            label = EXERCISE_LABELS.get(b.kind, b.kind.title())
            p = doc.add_paragraph()
            p.add_run(f"{label}{': ' + b.title if b.title else ''} ").bold = True
            _add_runs(p, b.runs)
        elif isinstance(b, FootnoteBlock):
            p = doc.add_paragraph()
            p.add_run(b.label + " ").bold = True
            _add_runs(p, b.runs)

    def rec(sections: list[Section], depth: int) -> None:
        for s in sections:
            if not s.enabled or (s.kind in GENERATED_KINDS and not s.blocks):
                continue
            if s.title and s.show_title:
                doc.add_heading(s.title, level=min(depth, 3))
            for b in s.blocks:
                block(b)
            rec(s.children, depth + 1)

    rec(book.sections, 1)
    notes.append("Styles and page layout are simplified in DOCX export; cover and generated pages are omitted")
    doc.save(str(path))
    return sorted(set(notes))
